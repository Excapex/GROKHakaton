"""Apply ChangeSet patches to local copies. Originals are never overwritten.

CAD / DWG is a design_task, not a patch. Stale base_hashes refuse the apply.
"""
from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from engine.changeset.plan import StaleSourceError


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return "sha256:" + h.hexdigest()


def apply_changeset(
    change_set: dict,
    sources: dict[str, Path],
    out_dir: Path,
) -> dict:
    """Copy sources to out_dir, then patch. sources keys are document_id."""
    current = {doc_id: file_hash(path) for doc_id, path in sources.items()}
    for doc_id, expected in change_set["base_hashes"].items():
        got = current.get(doc_id)
        if got != expected:
            raise StaleSourceError(f"{doc_id}: base {expected} != current {got}")

    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []
    copies: dict[str, Path] = {}
    for doc_id, src in sources.items():
        dest = out_dir / f"{doc_id}{src.suffix}"
        shutil.copy2(src, dest)
        copies[doc_id] = dest

    for patch in change_set.get("patches", []):
        fmt = patch["format"]
        dest = copies[patch["document_id"]]
        if fmt == "docx":
            _patch_docx(dest, patch["from"], patch["to"])
        elif fmt == "xlsx":
            _patch_xlsx(dest, patch["locator"], patch["from"], patch["to"])
        else:
            raise ValueError(f"unsupported patch format {fmt}")
        written.append(str(dest))

    skipped = [
        t for t in change_set.get("design_tasks", []) if t.get("reason") == "unsupported_format"
    ]
    provenance = {
        "schema_version": "1.0.0",
        "change_set_id": change_set.get("id"),
        "lifecycle_note": "applied_on_copies",
        "not_consent": True,
        "originals_untouched": True,
        "written": written,
        "design_tasks": skipped,
        "base_hashes": change_set.get("base_hashes", {}),
    }
    (out_dir / "provenance.json").write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return {
        "ok": True,
        "written": written,
        "design_tasks": skipped,
        "out_dir": str(out_dir),
        "provenance": str(out_dir / "provenance.json"),
    }


def _patch_docx(path: Path, old: str, new: str) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("Nedostaje python-docx. pip install -r requirements.txt") from exc
    doc = Document(str(path))
    hits = 0
    for p in doc.paragraphs:
        if old in p.text:
            for run in p.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    hits += 1
    if hits == 0:
        raise ValueError(f"docx: nije nađen tekst {old!r}")
    doc.save(str(path))


def _patch_xlsx(path: Path, locator: str, old: str, new: str) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("Nedostaje openpyxl. pip install -r requirements.txt") from exc
    # locator: Sheet1!C12
    if "!" not in locator:
        raise ValueError(f"xlsx locator mora biti Sheet!Cell, dobijeno {locator}")
    sheet_name, cell = locator.split("!", 1)
    wb = load_workbook(str(path))
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"nema lista {sheet_name}")
    ws = wb[sheet_name]
    value = ws[cell].value
    text = "" if value is None else str(value)
    if old not in text:
        raise ValueError(f"xlsx {locator}: nije nađen {old!r}")
    ws[cell] = text.replace(old, new)
    wb.save(str(path))
