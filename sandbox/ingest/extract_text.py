#!/usr/bin/env python3
"""Extract page/sheet text. page_no is physical index (1-based). Never invents pages."""
from __future__ import annotations

import json
import sys
from pathlib import Path

FULL_TEXT_MIN_CHARS = 40


def _readability(text: str) -> str:
    return "full" if len(text.strip()) >= FULL_TEXT_MIN_CHARS else "partial"


def extract_pdf(path: Path) -> list[dict]:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore[no-redef]
    doc = fitz.open(path)
    try:
        pages = []
        for i, page in enumerate(doc, start=1):
            text = page.get_text() or ""
            pages.append(
                {
                    "page_no": i,
                    "text": text,
                    "readability": _readability(text),
                    "char_count": len(text.strip()),
                }
            )
        return pages
    finally:
        doc.close()


def extract_docx(path: Path) -> list[dict]:
    from docx import Document

    text = "\n".join(p.text for p in Document(path).paragraphs)
    return [
        {
            "page_no": 1,
            "text": text,
            "readability": _readability(text),
            "char_count": len(text.strip()),
        }
    ]


def extract_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    pages = []
    try:
        for i, name in enumerate(book.sheetnames, start=1):
            sheet = book[name]
            lines = [name]
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join("" if cell is None else str(cell) for cell in row))
            text = "\n".join(lines)
            pages.append(
                {
                    "page_no": i,
                    "text": text,
                    "readability": _readability(text),
                    "char_count": len(text.strip()),
                }
            )
    finally:
        book.close()
    return pages


def extract(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = extract_pdf(path)
    elif suffix == ".docx":
        pages = extract_docx(path)
    elif suffix == ".xlsx":
        pages = extract_xlsx(path)
    else:
        pages = []
    return {"schema_version": "1.0.0", "page_count": len(pages), "pages": pages}


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: extract_text.py <file>", file=sys.stderr)
        return 2
    path = Path(sys.argv[1]).expanduser()
    if not path.is_file():
        print(f"nema fajla: {path}", file=sys.stderr)
        return 2
    print(json.dumps(extract(path), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
