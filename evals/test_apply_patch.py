#!/usr/bin/env python3
"""S11/S12-engine: apply patches to copies; refuse stale hash; never patch DWG."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine.changeset.plan import StaleSourceError, plan_r1  # noqa: E402
from sandbox.compute.apply import apply_changeset, file_hash  # noqa: E402


def fail(msg: str) -> None:
    print(f"FAIL {msg}", file=sys.stderr)
    raise SystemExit(1)


def _docx(path: Path, text: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def _xlsx(path: Path, value: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["C12"] = value
    wb.save(str(path))


def main() -> int:
    try:
        from docx import Document  # noqa: F401
        from openpyxl import Workbook  # noqa: F401
    except ImportError:
        print("evals/test_apply_patch: SKIP (pip install python-docx openpyxl)")
        return 0

    with tempfile.TemporaryDirectory() as raw:
        tmp = Path(raw)
        docx = tmp / "opis.docx"
        xlsx = tmp / "predmer.xlsx"
        _docx(docx, "Otpornost F60 za vrata.")
        _xlsx(xlsx, "F60")
        h_doc = file_hash(docx)
        h_xls = file_hash(xlsx)
        cs = plan_r1(
            gpzop_id="doc_opis",
            gpzop_hash=h_doc,
            predmer_id="doc_predmer",
            predmer_hash=h_xls,
            dwg_id="doc_cad",
        )
        out = tmp / "rev2"
        result = apply_changeset(cs, {"doc_opis": docx, "doc_predmer": xlsx}, out)
        if not result["ok"] or len(result["written"]) != 2:
            fail("apply wrote patches")
        if not any(t["reason"] == "unsupported_format" for t in result["design_tasks"]):
            fail("CAD remains design_task")
        from docx import Document

        text = "\n".join(p.text for p in Document(str(out / "doc_opis.docx")).paragraphs)
        if "F60" in text or "EI 60" not in text:
            fail("docx replace")
        from openpyxl import load_workbook

        val = str(load_workbook(out / "doc_predmer.xlsx")["Sheet1"]["C12"].value)
        if val != "EI 60":
            fail("xlsx cell")
        # original untouched
        orig = "\n".join(p.text for p in Document(str(docx)).paragraphs)
        if "F60" not in orig:
            fail("original must stay")
        try:
            stale = dict(cs)
            stale["base_hashes"] = {**cs["base_hashes"], "doc_opis": "sha256:stale"}
            apply_changeset(stale, {"doc_opis": docx, "doc_predmer": xlsx}, tmp / "nope")
            fail("stale hash must reject")
        except StaleSourceError:
            pass
        import json

        prov = json.loads((out / "provenance.json").read_text(encoding="utf-8"))
        if prov.get("originals_untouched") is not True or not prov.get("not_consent"):
            fail("provenance")
    print("evals/test_apply_patch: OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
