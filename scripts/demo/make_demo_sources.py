#!/usr/bin/env python3
"""Build the anon DOCX/XLSX originals the S12 apply demo patches.

Content mirrors `evals/fixtures/extract`; no client project data. Output is
gitignored because a DOCX/XLSX is a zip and zips do not belong in the repo.

  .venv/bin/python scripts/demo/make_demo_sources.py
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "sandbox" / "artifacts" / "source"

GPZOP_ID = "doc_anon_gpzop"
PREDMER_ID = "doc_anon_predmer"


def build_gpzop(path: Path) -> None:
    doc = Document()
    doc.add_heading("Glavni projekat zaštite od požara — anon primer", level=1)
    doc.add_paragraph("Objekat: anon fixture, sala 180 m2 za 90 lica.")
    doc.add_paragraph("Otpornost: F60 za protivpožarna vrata D-12.")
    doc.add_paragraph("Fasada: A1 klasa obloga izolacije.")
    doc.add_paragraph("Sigurnosna rasveta: autonomija 1 h.")
    doc.save(path)


def build_predmer(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Poz."
    sheet["B1"] = "Opis"
    sheet["C1"] = "Oznaka"
    sheet["A12"] = "1.12"
    sheet["B12"] = "Protivpožarna vrata D-12"
    sheet["C12"] = "F60"
    book.save(path)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    gpzop = OUT / "anon-gpzop.docx"
    predmer = OUT / "anon-predmer.xlsx"
    build_gpzop(gpzop)
    build_predmer(predmer)

    import sys

    sys.path.insert(0, str(ROOT))
    from sandbox.compute.apply import file_hash

    change_set = {
        "schema_version": "1.0.0",
        "id": "cs_r1_fire_mark",
        "lifecycle": "proposed",
        "base_hashes": {GPZOP_ID: file_hash(gpzop), PREDMER_ID: file_hash(predmer)},
        "patches": [
            {
                "document_id": GPZOP_ID,
                "format": "docx",
                "op": "replace_text",
                "locator": "body:fire-resistance-mark",
                "from": "F60",
                "to": "EI 60 prema SRPS EN 13501-2",
            },
            {
                "document_id": PREDMER_ID,
                "format": "xlsx",
                "op": "set_cell",
                "locator": "Sheet1!C12",
                "from": "F60",
                "to": "EI 60",
            },
        ],
        "design_tasks": [
            {
                "id": "dt_cad_mark",
                "description": "Oznaku otpornosti na DWG/DWFX uskladiti ručno; parser ne patch-uje CAD.",
                "reason": "unsupported_format",
                "document_id": "doc_anon_cad",
            }
        ],
        "dependencies": ["patch:docx-opis", "patch:xlsx-predmer"],
        "approval": {"state": "proposed"},
        "artifact_ids": [],
    }
    plan = OUT / "changeset.json"
    plan.write_text(json.dumps(change_set, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(json.dumps({"gpzop": str(gpzop), "predmer": str(predmer), "change_set": str(plan)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
